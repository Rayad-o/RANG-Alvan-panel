"""تبدیل فایل‌های گزارش اکسل (XML Spreadsheet 2003 با پسوند .xls) به .xlsx واقعی.

هر فایل خروجی:
  • سرستون واقعی در سطر ۱ (بدون سطر عنوان/جمع‌کل/خالی) با حروف فارسی ی/ک
  • فقط ستون‌هایی که سایت می‌خواند، با نام دقیق مورد انتظار endpoint مربوط
  • ستون‌های عددی به‌صورت عدد و کدها/شماره‌ها/تاریخ‌ها به‌صورت متن
    (تا صفرهای جلو حفظ شوند، مثلاً کد بانک «039»)

خروجی: پوشهٔ exels/converted/

 اجرا:  python convert_exels.py
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
SRC_DIR = ROOT / "exels"
DST_DIR = SRC_DIR / "converted"

SS_NS = "urn:schemas-microsoft-com:office:spreadsheet"
_SS = {"ss": SS_NS}

# حروف عربی → فارسی (همان نرمال‌سازی backend/app/core/excel.py)
def normalize(text: str) -> str:
    return (
        text.replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace("\u00a0", " ")
        .strip()
    )


def read_rows(path: Path) -> list[list[str | None]]:
    """خواندن SpreadsheetML با رعایت ss:Index (ستون‌های جاافتاده پر می‌شوند)."""
    content = path.read_bytes()
    # بعضی فایل‌ها BOM دارند؛ ET خودش نمی‌خواند
    if content.startswith(b"\xef\xbb\xbf"):
        content = content[3:]
    root = ET.fromstring(content)

    table = None
    for worksheet in root.findall(".//ss:Worksheet", _SS):
        table = worksheet.find("ss:Table", _SS)
        if table is not None:
            break
    if table is None:
        raise RuntimeError(f"جدولی در {path.name} پیدا نشد")

    rows: list[list[str | None]] = []
    for row in table.findall("ss:Row", _SS):
        values: list[str | None] = []
        column = 0
        for cell in row.findall("ss:Cell", _SS):
            raw_index = cell.get(f"{{{SS_NS}}}Index")
            if raw_index is not None:
                try:
                    column = int(raw_index) - 1
                except ValueError:
                    column = len(values)
            while len(values) < column:
                values.append(None)
            data_el = cell.find("ss:Data", _SS)
            values.append(data_el.text if data_el is not None else None)
            column += 1
        rows.append(values)
    return rows


def is_empty(value) -> bool:
    return value is None or str(value).strip() == ""


def find_header(rows: list[list[str | None]]) -> int:
    """اولین سطری که بیش از یک سلول پر دارد (سطرهای عنوان تک‌سلولی رد می‌شوند)."""
    for i, row in enumerate(rows):
        if sum(1 for v in row if not is_empty(v)) > 1:
            return i
    raise RuntimeError("سطر سرستون پیدا نشد")


NUM_RE = re.compile(r"^-?\d+(?:\.\d+)?$")


def to_number(text: str):
    """متن عددی → int/float؛ در غیر این صورت همان متن برمی‌گردد."""
    cleaned = text.replace(",", "").replace("٬", "").replace("\u200c", "").strip()
    if not NUM_RE.match(cleaned):
        return text
    if "." in cleaned:
        return float(cleaned)
    try:
        return int(cleaned)
    except ValueError:
        return text


# ---------------------------------------------------------------------------
# نگاشت هر فایل: ستون مبدأ (بعد از نرمال‌سازی) → ستون خروجی + نوع مقدار
#   t = متن (کد، شماره، تاریخ، نام)   n = عدد
#   key = ستون‌هایی که خالی‌بودنشان یعنی سطر جمع‌کل/خلاصه است و حذف می‌شود
# ---------------------------------------------------------------------------
SPECS: list[dict] = [
    {
        "src": "مرور فروش.xls",
        "out": "مرور فروش.xlsx",
        "cols": [("کد", "کد", "t"), ("عنوان", "عنوان", "t"), ("مانده", "مانده", "n")],
        "key": ["کد", "عنوان"],
    },
    {
        "src": "مرور تامين كنندگان.xls",
        "out": "مرور تامين كنندگان.xlsx",
        "cols": [
            ("کد", "کد", "t"),
            ("عنوان", "عنوان", "t"),
            ("گردش بدهکار", "گردش بدهکار", "n"),
            ("مانده بدهکار", "مانده بدهکار", "n"),
            ("مانده بستانکار", "مانده بستانکار", "n"),
        ],
        "key": ["کد", "عنوان"],
    },
    {
        "src": "مرور_مانده_ساير_حسابهاي_پرداختني_و_پر_ژه.xls",
        "out": "مرور_مانده_ساير_حسابهاي_پرداختني_و_پر_ژه.xlsx",
        "cols": [
            ("کد", "کد", "t"),
            ("عنوان", "عنوان", "t"),
            ("گردش بدهکار", "گردش بدهکار", "n"),
            ("مانده بدهکار", "مانده بدهکار", "n"),
            ("مانده بستانکار", "مانده بستانکار", "n"),
        ],
        "key": ["کد", "عنوان"],
    },
    {
        "src": "مرور تعدادي انبار محصول.xls",
        "out": "مرور تعدادي انبار محصول.xlsx",
        "cols": [
            ("کد کالا", "کد کالا", "t"),
            ("عنوان کالا", "عنوان کالا", "t"),
            ("موجودی", "موجودی", "n"),
        ],
        "key": ["کد کالا", "عنوان کالا"],
    },
    {
        "src": "مرور مبلغي انبار محصول.xls",
        "out": "مرور مبلغي انبار محصول.xlsx",
        "cols": [
            ("کد کالا", "کد کالا", "t"),
            ("عنوان کالا", "عنوان کالا", "t"),
            ("موجودی-مقدار", "موجودی", "n"),
            ("موجودی-مبلغ", "موجودی-مبلغ", "n"),
            ("موجودی-فی متوسط", "موجودی-فی متوسط", "n"),
        ],
        "key": ["کد کالا", "عنوان کالا"],
    },
    {
        "src": "مرور بانكها دريافت و پرداخت.xls",
        "out": "مرور بانكها دريافت و پرداخت.xlsx",
        "cols": [
            ("کد تفصیلی", "کد تفصیلی", "t"),
            ("عنوان", "عنوان", "t"),
            ("اول دوره", "اول دوره", "n"),
            ("دریافت ها", "دریافت ها", "n"),
            ("پرداخت ها", "پرداخت ها", "n"),
            ("مانده", "مانده", "n"),
        ],
        "key": ["کد تفصیلی", "عنوان"],
    },
    {
        "src": "مرور دريافت بتفكيك نوع و تاريخ.xls",
        "out": "مرور دريافت بتفكيك نوع و تاريخ.xlsx",
        "cols": [
            ("نوع دریافت", "نوع دریافت", "t"),
            ("طرف مقابل", "طرف مقابل", "t"),
            ("شماره", "شماره", "t"),
            ("تاریخ", "تاریخ", "t"),
            ("مبلغ دریافت", "مبلغ دریافت", "n"),
            ("شرح", "شرح", "t"),
        ],
        "key": ["طرف مقابل"],
    },
    {
        "src": "گزارش پرداخت ب تفكيك نوع و روز.xls",
        "out": "گزارش پرداخت ب تفكيك نوع و روز.xlsx",
        "cols": [
            ("نوع پرداخت", "نوع پرداخت", "t"),
            ("طرف مقابل", "طرف مقابل", "t"),
            ("شماره", "شماره", "t"),
            ("تاریخ", "تاریخ", "t"),
            ("اقلام", "اقلام", "t"),
            ("مبلغ پرداخت", "مبلغ پرداخت", "n"),
            ("شرح", "شرح", "t"),
        ],
        "key": ["طرف مقابل"],
    },
    {
        # چک‌های پرداختنی در انتظار (بدون ستون مبلغ در فایل مبدأ)
        "src": "مرور بانك چك پرداختني مهر.xls",
        "out": "مرور بانك چك پرداختني مهر.xlsx",
        "cols": [
            ("تفصیلی", "تفصیلی", "t"),
            ("طرف مقابل", "طرف مقابل", "t"),
            ("تاریخ سررسید", "تاریخ سررسید", "t"),
            ("شماره اعلامیه", "شماره اعلامیه", "t"),
            ("تاریخ اعلامیه", "تاریخ اعلامیه", "t"),
            ("تاریخ وصول", "تاریخ وصول", "t"),
            ("حساب بانکی", "حساب بانکی", "t"),
            ("نوع", "نوع", "t"),
        ],
        "key": ["تفصیلی", "طرف مقابل"],
    },
    {
        "src": "مرور چكهاي پرداختني پاس شده.xls",
        "out": "مرور چكهاي پرداختني پاس شده.xlsx",
        "cols": [
            ("تفصیلی", "تفصیلی", "t"),
            ("طرف مقابل", "طرف مقابل", "t"),
            ("مبلغ", "مبلغ", "n"),
            ("نرخ", "نرخ", "n"),
            ("تاریخ سررسید", "تاریخ سررسید", "t"),
            ("شماره اعلامیه", "شماره اعلامیه", "t"),
            ("تاریخ اعلامیه", "تاریخ اعلامیه", "t"),
            ("تاریخ وصول", "تاریخ وصول", "t"),
            ("حساب بانکی", "حساب بانکی", "t"),
            ("نوع", "نوع", "t"),
        ],
        "key": ["تفصیلی", "طرف مقابل"],
    },
    {
        "src": "مرور_بانك_چك_هاي_دريافتني_شهريور_واگزار_به_بانك.xls",
        "out": "مرور_بانك_چك_هاي_دريافتني_شهريور_واگزار_به_بانك.xlsx",
        "cols": [
            ("تفصیلی", "تفصیلی", "t"),
            ("طرف مقابل", "طرف مقابل", "t"),
            ("نوع", "نوع", "t"),
            ("شماره", "شماره", "t"),
            ("مبلغ", "مبلغ", "n"),
            ("تاریخ سررسید", "تاریخ سررسید", "t"),
            ("تاریخ رسید", "تاریخ رسید", "t"),
            ("بانک", "بانک", "t"),
            ("وضعیت", "وضعیت", "t"),
            ("تاریخ وصول، واخواست", "تاریخ وصول، واخواست", "t"),
            ("صاحب چک", "صاحب چک", "t"),
            ("شرح", "شرح", "t"),
        ],
        "key": ["طرف مقابل"],
    },
    {
        "src": "مرور_چك_دريافتي_خرج_شده_مهرماه_سررسيد.xls",
        "out": "مرور_چك_دريافتي_خرج_شده_مهرماه_سررسيد.xlsx",
        "cols": [
            ("تفصیلی", "تفصیلی", "t"),
            ("طرف مقابل", "طرف مقابل", "t"),
            ("نوع", "نوع", "t"),
            ("شماره", "شماره", "t"),
            ("مبلغ", "مبلغ", "n"),
            ("تاریخ سررسید", "تاریخ سررسید", "t"),
            ("تاریخ رسید", "تاریخ رسید", "t"),
            ("بانک", "بانک", "t"),
            ("وضعیت", "وضعیت", "t"),
            ("تاریخ وصول، واخواست", "تاریخ وصول، واخواست", "t"),
            ("صاحب چک", "صاحب چک", "t"),
            ("شرح", "شرح", "t"),
        ],
        "key": ["طرف مقابل"],
    },
]


def convert(spec: dict) -> None:
    src_path = SRC_DIR / spec["src"]
    if not src_path.exists():
        raise FileNotFoundError(src_path)

    rows = read_rows(src_path)
    header_idx = find_header(rows)
    header_cells = [normalize(str(v)) if v is not None else "" for v in rows[header_idx]]
    col_index = {name: i for i, name in enumerate(header_cells) if name}

    missing = [src for src, _, _ in spec["cols"] if src not in col_index]
    if missing:
        raise RuntimeError(f"{spec['src']}: ستون مبدأ پیدا نشد: {missing} | سرستون‌ها: {header_cells}")

    key_idx = [col_index[k] for k in spec["key"]]

    records: list[list] = []
    dropped = 0
    for row in rows[header_idx + 1:]:
        # سطر کاملاً خالی
        if all(is_empty(v) for v in row):
            dropped += 1
            continue
        # سطر جمع‌کل/خلاصه: یکی از ستون‌های کلیدی خالی است
        if any(i >= len(row) or is_empty(row[i]) for i in key_idx):
            dropped += 1
            continue

        record = []
        for src, _, kind in spec["cols"]:
            value = row[col_index[src]] if col_index[src] < len(row) else None
            if is_empty(value):
                record.append(None)
            elif kind == "n":
                record.append(to_number(str(value)))
            else:
                record.append(str(value).strip())
        records.append(record)

    # نوشتن xlsx: سرستون در سطر ۱
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append([out for _, out, _ in spec["cols"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for record in records:
        ws.append(record)

    # عرض ستون‌ها بر اساس محتوا
    for i, (_, out, _) in enumerate(spec["cols"], start=1):
        width = max([len(str(out)) + 2] + [len(str(r[i - 1])) + 2 for r in records if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 60)

    dst_path = DST_DIR / spec["out"]
    wb.save(dst_path)
    print(f"  ✓ {spec['src']}  →  {dst_path.relative_to(ROOT)}  ({len(records)} سطر، {dropped} سطر حذف‌شده)")


def main() -> None:
    DST_DIR.mkdir(parents=True, exist_ok=True)
    print(f"تبدیل {len(SPECS)} فایل از {SRC_DIR} به {DST_DIR}:\n")
    for spec in SPECS:
        convert(spec)
    print("\nهمهٔ فایل‌ها با موفقیت تبدیل شدند.")


if __name__ == "__main__":
    main()
